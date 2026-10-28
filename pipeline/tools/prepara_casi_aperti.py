"""Foglio dei casi che l'estrazione non sa decidere da sola: pipeline/CASI_APERTI.xlsx.

Uso: python3 tools/prepara_casi_aperti.py
Si apre il file, si guarda la pagina indicata e si scrive la scelta nella colonna DECISIONE
(menu a tendina). esporta_dataset.py legge il foglio e applica le decisioni a ogni esportazione.
Rigenerare il foglio NON cancella le decisioni gia' prese: vengono riportate per chiave."""
import os, re, difflib, unicodedata
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

QUI = os.path.dirname(os.path.abspath(__file__)); PIPE = os.path.dirname(QUI); EST = os.path.join(PIPE, 'extraction')
OUT = os.path.join(PIPE, 'CASI_APERTI.xlsx')
nomi = pd.read_csv(os.path.join(QUI, 'costituenti_nomi.csv'))
nomi['fine_mandato'] = nomi.fine_mandato.fillna('1948-01-31')
nomi['etichetta'] = nomi.cognome + ' ' + nomi.nome
ETI = dict(zip(nomi.id_dep, nomi.etichetta))
def norm(s): return re.sub(r'[^A-Z]', '', unicodedata.normalize('NFKD', str(s)).encode('ascii', 'ignore').decode().upper())
nomi['k'] = nomi.cognome.map(norm)
def in_carica(d): return nomi[(nomi.inizio_mandato <= d) & (d <= nomi.fine_mandato)] if isinstance(d, str) else nomi
def simili(t, d, n=6):
    c = in_carica(d); k = norm(t)
    best = difflib.get_close_matches(k, list(c.k), n=n, cutoff=0.5)
    return [e for b in best for e in c[c.k == b].etichetta][:n]
URL = 'https://documenti.camera.it/_dati/costituente/lavori/Assemblea/{s}/{s}nc.pdf#page={p}'
NESSUNO = 'nessuno (rumore OCR / non è un costituente)'

fogli = {}
# ---------- 1. nomi non risolti negli emendamenti ----------
T = pd.read_csv(os.path.join(EST, 'menzioni_tutte.csv')); F = pd.read_csv(os.path.join(EST, 'menzioni_nomi.csv'))
E = pd.read_csv(os.path.join(EST, 'emendamenti.csv')).set_index('id_emendamento')
u = F[F.abbinamento.isin(['ambiguo', 'non trovato'])].merge(T, on='id_menzione', suffixes=('', '_m'))
righe = []
for r in u.itertuples():
    lettere = re.sub(r'[^A-Za-zÀ-ÿ]', '', str(r.nome_stampato))
    if len(lettere) < 4 or re.search(r'Ministr|Presidente|Relatore', str(r.nome_stampato)): continue
    cand = [ETI[c] for c in str(r.candidati).split('|') if c in ETI] if isinstance(r.candidati, str) else simili(r.nome_stampato, r.data)
    if pd.notna(r.id_emendamento) and r.id_emendamento in E.index:
        e = E.loc[r.id_emendamento]; fs = str(e.firmatari).split('|')
        if any(c in fs for c in str(r.candidati).split('|')): continue
        chiave_em = f'{e.prima_seduta}|{e.pagina}|{e.firmatari}'
        cosa = f"Il nome manca nell'emendamento già nel dataset ({e.n_firmatari} firmatari: {', '.join(ETI.get(x, x) for x in fs)[:200]}). Se è un costituente, sceglilo: verrà aggiunto all'elenco."
        opz = cand + [NESSUNO]
    else:
        chiave_em = ''
        cosa = "Nessun emendamento nel dataset: l'unico nome stampato non è stato risolto. Se sulla pagina è davvero presentato un emendamento, scegli chi lo presenta: verrà creato a firma singola."
        opz = cand + [NESSUNO, 'non è un emendamento']
    righe.append(dict(caso=f'E{len(righe) + 1:03d}', seduta=r.seduta, pagina=r.pagina, link_pagina=URL.format(s=r.seduta, p=r.pagina),
                      file_locale=f'pipeline/pdf/{r.seduta}.pdf, pagina {r.pagina}', nome_stampato=r.nome_stampato,
                      elenco_stampato=r.testo_stampato, cosa_fare=cosa, opzioni=' | '.join(opz), DECISIONE='', nota='',
                      contesto_ocr=str(r.contesto)[:600], chiave=f'{r.seduta}|{r.pagina}|{r.nome_stampato}', chiave_emendamento=chiave_em))
fogli['emendamenti_nomi'] = pd.DataFrame(righe)

# ---------- 2. elenchi chiusi da «ed altri» ----------
D = pd.read_csv(os.path.join(os.path.dirname(PIPE), 'data', 'it', 'emendamenti_nodi.csv'))
righe = []
for r in D[D.completo == 0].itertuples():
    fs = str(r.firmatari_id).split('|')
    righe.append(dict(caso=f'A{len(righe) + 1:03d}', seduta=r.prima_seduta, pagina=r.pagina, link_pagina=r.url_pdf,
                      file_locale=f'pipeline/pdf/{r.prima_seduta}.pdf, pagina {r.pagina}',
                      firmatari_attuali=', '.join(ETI.get(x, x) for x in fs),
                      cosa_fare="L'elenco è chiuso da «ed altri». Cerca nella stessa seduta o nelle vicine se lo stesso emendamento è stampato con tutte le firme (spesso il Presidente lo rilegge prima del voto). Se le trovi, scrivi in DECISIONE i nomi mancanti separati da punto e virgola, come «Rossi Paolo; Corsi»; altrimenti scegli «nessun elenco completo stampato».",
                      opzioni='nessun elenco completo stampato | <nomi separati da ;>', DECISIONE='', nota='',
                      testo=str(r.testo)[:300], chiave=f'{r.prima_seduta}|{r.pagina}|{r.firmatari_id}'))
fogli['emendamenti_ed_altri'] = pd.DataFrame(righe)

# ---------- 3. voti: nomi non abbinati o ambigui ----------
V = pd.read_csv(os.path.join(EST, 'voti.csv')); N = pd.read_csv(os.path.join(EST, 'voti_nominali.csv'), keep_default_na=False)
data_v = dict(zip(V.id_voto, V.data)); pag_v = dict(zip(V.id_voto, V.pagina)); sed_v = dict(zip(V.id_voto, V.seduta))
STOP = set("""DICHIARO CHIUSA VOTAZIONE VOTAZIONI SEGRETARI SEGRETARIO COMPUTO VOTI VOTO ONOREVOLI INVITO PRESIDENTE PROCEDERE
RISULTATO NOMINALE SCRUTINIO APPELLO COMUNICO ANNUNCIO SEGUENTE PRESENTI VOTANTI MAGGIORANZA ASSEMBLEA SEDUTA CAMERA DEPUTATI
DEPUTATO CONGEDO ASTENUTI SONO HANNO RISPOSTO RISPONDONO ARTICOLO EMENDAMENTO PRESIDENZA RELATORE MINISTRO FARE FANNO PROCEDONO
ERRATA CORRIGE VEDI ANTIMERIDIANA""".split())
PROPRI = {t for x in nomi.nome for t in norm(x + ' ').split()} | {norm(t) for x in nomi.nome for t in str(x).split()}
k_di = dict(zip(nomi.id_dep, nomi.k))
righe = []
for (vid, voto), g in N.groupby(['id_voto', 'voto'], sort=False):
    g = g.reset_index(drop=True)
    for p, r in g.iterrows():
        if r.abbinamento not in ('non trovato', 'ambiguo'): continue
        t = norm(r.nome_stampato)
        if r.abbinamento == 'non trovato' and (len(t) < 4 or t in STOP or t in PROPRI): continue
        lo = next((g.nome_stampato[j] for j in range(p - 1, -1, -1) if g.id_dep[j]), '')
        hi = next((g.nome_stampato[j] for j in range(p + 1, len(g)) if g.id_dep[j]), '')
        klo = next((k_di.get(g.id_dep[j], '') for j in range(p - 1, -1, -1) if g.id_dep[j]), '')
        khi = next((k_di.get(g.id_dep[j], '') for j in range(p + 1, len(g)) if g.id_dep[j]), '')
        c = in_carica(data_v.get(vid))
        presenti = set(N[(N.id_voto == vid) & (N.id_dep != '')].id_dep)
        fascia = c[(c.k >= klo) & (c.k <= (khi or 'ZZZ')) & ~c.id_dep.isin(presenti)] if klo else c.iloc[0:0]
        opz = list(fascia.etichetta)[:15] if 0 < len(fascia) <= 15 else simili(r.nome_stampato, data_v.get(vid))
        righe.append(dict(caso=f'V{len(righe) + 1:03d}', id_voto=f'v{vid}', voto=voto, seduta=sed_v[vid],
                          link_pagina=URL.format(s=sed_v[vid], p=pag_v[vid]),
                          file_locale=f'pipeline/pdf/{sed_v[vid]}.pdf, dalla pagina {pag_v[vid]} (elenco del «{voto}»)',
                          nome_stampato=r.nome_stampato, tra=f'{lo} … {hi}',
                          cosa_fare="Trova il nome nell'elenco stampato del voto indicato (è in ordine alfabetico, fra i due nomi della colonna «tra»). Scegli il costituente; «nessuno» se è una parola del verbale o un pezzo di un nome già contato.",
                          opzioni=' | '.join(opz + [NESSUNO]), DECISIONE='', nota='', chiave=f'{vid}|{voto}|{r.nome_stampato}|{p}'))
fogli['voti_nomi'] = pd.DataFrame(righe)

# ---------- 4. voti: abbinamenti approssimati ----------
righe = []
for (vid, voto), g in N.groupby(['id_voto', 'voto'], sort=False):
    g = g.reset_index(drop=True)
    for p, r in g.iterrows():
        if r.abbinamento not in ('approssimato', 'fascia alfabetica', 'omonimo risolto'): continue
        lo = next((g.nome_stampato[j] for j in range(p - 1, -1, -1) if g.id_dep[j]), '')
        hi = next((g.nome_stampato[j] for j in range(p + 1, len(g)) if g.id_dep[j]), '')
        righe.append(dict(caso=f'P{len(righe) + 1:03d}', id_voto=f'v{vid}', voto=voto, seduta=sed_v[vid],
                          link_pagina=URL.format(s=sed_v[vid], p=pag_v[vid]), nome_stampato=r.nome_stampato,
                          abbinato_a=ETI.get(r.id_dep, r.id_dep), metodo=r.abbinamento, tra=f'{lo} … {hi}',
                          cosa_fare="Controlla che il nome stampato sia proprio la persona abbinata. Lascia vuoto o scegli «giusto» se lo è; altrimenti scegli «sbagliato: scarta» o scrivi «Cognome Nome» della persona giusta.",
                          opzioni='giusto | sbagliato: scarta | <Cognome Nome>', DECISIONE='', nota='',
                          chiave=f'{vid}|{voto}|{r.nome_stampato}|{r.id_dep}'))
fogli['voti_approssimati'] = pd.DataFrame(righe)

# ---------- 5. casi incerti gia' decisi in modo provvisorio ----------
fogli['incerti'] = pd.DataFrame([
    dict(caso='I001', seduta='sed322', pagina=17, link_pagina=URL.format(s='sed322', p=17),
         cosa_fare="Elenco liberale e qualunquista (Martino Gaetano, Villabruna, Morelli Renato, Bonino, Colitto, Perrone Capano, Condorelli, Colonna di Paliano, «Massa», Rodinò, Cifaldi). «Massa» non è un costituente: oggi è attribuito a Mazza Crescenzo (UQ). Guarda la stampa: se si legge Mazza, conferma; se si legge Massola (PCI, improbabile in questo elenco) sceglilo; se il nome è illeggibile, «nessuno».",
         opzioni='Mazza Crescenzo | Massola Umberto Paolo | nessuno', DECISIONE='', nota='', chiave='sed322_massa'),
    dict(caso='I002', seduta='sed104/sed105', pagina='sed104 p23, sed105 p8', link_pagina=URL.format(s='sed105', p=8),
         cosa_fare="Emendamento Martino. In sed104 p23 Martino lo dice presentato «dagli onorevoli Labriola, Della Seta, Lucifero, Corbino e da me» (5 nomi); in sed105 p8 il Presidente lo ristampa con anche Caronia (6 nomi). Oggi si tiene l'elenco a 6. Guarda le due pagine e scegli quale elenco è quello dell'emendamento.",
         opzioni='elenco a 6 (con Caronia) | elenco a 5 (senza Caronia)', DECISIONE='', nota='', chiave='sed104_martino'),
])

# ---------- decisioni gia' prese: si riportano per chiave ----------
vecchie = {}
if os.path.exists(OUT):
    wb0 = load_workbook(OUT)
    for nome in wb0.sheetnames:
        ws = wb0[nome]; hdr = [c.value for c in ws[1]]
        if 'chiave' not in hdr or 'DECISIONE' not in hdr: continue
        ic, idd, inn = hdr.index('chiave'), hdr.index('DECISIONE'), hdr.index('nota')
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[idd] or row[inn]: vecchie[(nome, row[ic])] = (row[idd] or '', row[inn] or '')

wb = Workbook(); ws0 = wb.active; ws0.title = 'LEGGIMI'
testo = [
    'Casi aperti del dataset della Costituente', '',
    'Ogni foglio è un tipo di caso. Per ciascuna riga: apri «link_pagina» (o il PDF locale indicato), guarda la stampa,',
    'scegli nella colonna DECISIONE (gialla) una delle voci del menu a tendina, o scrivi il nome come indicato.',
    'Lascia vuota la DECISIONE per i casi che non vuoi decidere: restano come sono oggi.', '',
    'Regola: si sceglie solo ciò che è stampato. Se la pagina non permette di decidere, scrivi «nessuno» e una nota.', '',
    'Dopo aver salvato: python3 pipeline/esporta_dataset.py  (applica le decisioni e stampa cosa ha cambiato)', '',
    'Priorità, per quanto contano sui legami:',
    '  1. emendamenti_nomi: firmatari mancanti o emendamenti persi (creano o completano archi di co-firma)',
    '  2. incerti: due casi già decisi in modo provvisorio',
    '  3. voti_nomi: voti non attribuiti a nessuno',
    '  4. emendamenti_ed_altri: elenchi troncati, si completano solo se l\'elenco intero è stampato altrove',
    '  5. voti_approssimati: controllo di abbinamenti probabilmente giusti (già coerenti con l\'ordine alfabetico)',
]
for i, t in enumerate(testo, 1): ws0.cell(i, 1, t)
ws0['A1'].font = Font(bold=True, size=14); ws0.column_dimensions['A'].width = 120
giallo = PatternFill('solid', fgColor='FFF2CC')
for nome, df in fogli.items():
    ws = wb.create_sheet(nome)
    if df.empty: ws.cell(1, 1, 'nessun caso'); continue
    for (n, k), (d, nt) in vecchie.items():
        if n == nome: df.loc[df.chiave == k, ['DECISIONE', 'nota']] = [d, nt]
    cols = list(df.columns)
    for j, c in enumerate(cols, 1):
        ws.cell(1, j, c).font = Font(bold=True)
    for i, row in enumerate(df.itertuples(index=False), 2):
        for j, v in enumerate(row, 1):
            cell = ws.cell(i, j, v)
            if cols[j - 1] == 'link_pagina': cell.hyperlink = v; cell.font = Font(color='0563C1', underline='single')
            if cols[j - 1] == 'DECISIONE': cell.fill = giallo
            if cols[j - 1] in ('cosa_fare', 'contesto_ocr', 'elenco_stampato', 'opzioni', 'testo'): cell.alignment = Alignment(wrap_text=True, vertical='top')
        opz = [o.strip() for o in str(df.opzioni.iloc[i - 2]).split('|') if o.strip() and not o.strip().startswith('<')]
        lista = '"' + ','.join(o.replace(',', ' ').replace('"', '') for o in opz) + '"'
        if len(lista) < 255 and opz:
            dv = DataValidation(type='list', formula1=lista, allow_blank=True, showErrorMessage=False); ws.add_data_validation(dv)
            dv.add(ws.cell(i, cols.index('DECISIONE') + 1))
    for j, c in enumerate(cols, 1):
        ws.column_dimensions[ws.cell(1, j).column_letter].width = {'cosa_fare': 60, 'contesto_ocr': 60, 'opzioni': 45, 'DECISIONE': 32,
                                                                   'elenco_stampato': 40, 'link_pagina': 20, 'testo': 40, 'nota': 30}.get(c, 16)
    ws.freeze_panes = 'B2'
wb.save(OUT)
print('scritto', OUT, {k: len(v) for k, v in fogli.items()}, '| decisioni riportate:', len(vecchie))
